import model as M, model2 as M2, hier as HI
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
F="Arial"; thin=Side(style="thin",color="D5CCD0"); B=Border(left=thin,right=thin,top=thin,bottom=thin)
PL="7A2E4D"; HDR=PatternFill("solid",fgColor=PL); STEP=PatternFill("solid",fgColor="EAD7DF"); BAND=[PatternFill("solid",fgColor="FFFFFF"),PatternFill("solid",fgColor="F7F3F5")]
STF={"Saling melengkapi":"D9EAD3","Unik":"EEEEEE","Identik":"D9EAD3","Identik (disalin)":"D9EAD3","Kode sama, isi beda":"FCE8B2","Maksud sama, kode/struktur beda":"F4CCCC","Nama sama, kode beda":"F4CCCC"}
SHEET={t:(t if t!="RJ" else "Rawat Jalan")[:31] for t,_ in M.TITLES}
LSHEET="Daftar Pilihan & Lampiran"
wb=Workbook(); links=[]  # (sheet, cell, target_sheet, target_row)
pos={}  # var id -> (sheet,row) ; group id -> row in Konsolidasi ; ringkasan row
def hdr(ws,r,H,W):
    for i,h in enumerate(H,1):
        c=ws.cell(r,i,h); c.font=Font(name=F,bold=True,color="FFFFFF",size=10); c.fill=HDR; c.alignment=Alignment(wrap_text=True,vertical="center"); c.border=B
        ws.column_dimensions[get_column_letter(i)].width=W[i-1]
def cell(ws,r,c,v,bold=False,color="000000",fill=None,size=9,wrap=True):
    x=ws.cell(r,c,v if v not in ("",None) else None); x.font=Font(name=F,size=size,bold=bold,color=color); x.alignment=Alignment(wrap_text=wrap,vertical="top"); x.border=B
    if fill: x.fill=fill
    return x
def vname(i): v=M.BYID[i]; return f"{i} {v['var']}"
# ---- Panduan
wp=wb.active; wp.title="Panduan"
lines=[("SATUSEHAT — semua use case: variabel, elemen FHIR, nilai, crosscheck & ringkasan resource",True),
 ("Sumber: "+"; ".join(M.SRC[t] for t,_ in M.TITLES)+". satusehat.kemkes.go.id",False),("",False),
 ("Sheet",True),("Satu sheet per use case — satu baris = satu elemen FHIR + nilainya (.code & .display terpisah; baris berulang = pilihan jawaban). * = elemen wajib.",False),
 ("Kolom 'Muncul di (konsolidasi)' → hyperlink ke kartu variabel di sheet Konsolidasi. Kolom 'Kode juga dipakai di' → variabel lain (judul mana pun) yang memakai kode yang sama.",False),
 ("Ringkasan Crosscheck — satu baris per konsep (deduplikasi lintas judul), status, dan daftar perbedaan. Konsolidasi — elemen gabungan per konsep dengan tanda ada di ANC / RJ.",False),
 ("Daftar Pilihan & Lampiran — semua pilihan jawaban (system–code–display–keterangan) dan lampiran kode, satu blok per daftar (ID P-xxx / nama lampiran); ditautkan dari baris elemen terkait.",False),("",False),("Status crosscheck",True),
 ("Unik — hanya ada di satu variabel.",False),("Identik / Identik (disalin) — elemen & nilai sama. 'Disalin' = tahap ANC/HIV yang di playbook merujuk ke modul Rawat Jalan, detailnya disalin dari RJ.",False),
 ("Saling melengkapi — kode sama (atau konsep dipasangkan), perbedaannya hanya elemen/pilihan tambahan di salah satu use case; tidak ada nilai yang bertentangan (contoh: interpretation di ANC + valueString di RJ).",False),
 ("Kode sama, isi beda — kode konsep utama sama, tetapi ada nilai yang bertentangan pada elemen yang sama (unit, category, kode pilihan). Coding tambahan WHO/Kemkes anc-custom-codes tidak dihitung.",False),
 ("Maksud sama, kode/struktur beda — konsepnya setara tetapi kode, category, atau resource berbeda → potensi inkonsistensi dokumentasi.",False),
 ("Elemen generik (subject, patient, encounter, performer, author, source, effectiveDateTime) tidak dipakai untuk menentukan kecocokan, tapi tetap ditampilkan.",False),
 ("Ringkasan Resource — per resource: elemen apa yang biasanya ada (Inti ≥80% variabel, Umum 40–79%, Kadang 10–39%, Jarang <10%), elemen pilihan [x] dihitung sekali (cukup salah satu tipe data). Sheet 'Resource - Nilai' merinci nilai/kode per path + kode tambahan dari Lampiran Standar Terminologi v10.3 (ditandai).",False),
 ("Kolom Induk/Peran/Level & baris '▼' — variabel yang sebenarnya satu kesatuan (item kuesioner bertingkat, komponen pengukuran, bagian dokumen Composition, kelompok playbook) ditampilkan sebagai induk–anak; sheet 'Struktur Induk-Anak' membandingkan induk yang sama antar use case.",False),
 ("Baris '(rujukan)' — tahap yang di playbook hanya merujuk ke modul lain (mis. Rawat Jalan/IGD/Rawat Inap) tanpa merinci elemen; ditampilkan sebagai judul tahap + tautan playbook, tidak disalin.",False),
 ("Baris berlatar biru muda dengan tanda PELENGKAP — elemen yang tidak ada di playbook ini tetapi ada di playbook lain dalam satu konsep berstatus 'Saling melengkapi'; sumbernya disebut di kolom Keterangan.",False),
 ("Deskripsi variabel — dari kalimat playbook yang menyebut variabel tsb; bila tidak ada, dibuat otomatis oleh Claude (ditandai 'Claude, belum diverifikasi').",False),
 ("Nilai seperti Patient/{…}, Encounter/{…} = referensi; nilai (status) / (intent) = elemen wajib yang nilainya tidak dirinci di playbook.",False)]
for i,(t,b) in enumerate(lines,1):
    c=wp.cell(i,1,t); c.font=Font(name=F,bold=b,size=13 if i==1 else 10); c.alignment=Alignment(wrap_text=True)
wp.column_dimensions["A"].width=150
from collections import Counter
r=len(lines)+2; wp.cell(r,1,"Jumlah konsep per status: "+" · ".join(f"{k}: {n}" for k,n in Counter(g['status'] for g in M.G).items())).font=Font(name=F,size=10,bold=True)
# ---- title sheets
H=["ID variabel","Tahap alur","Kelompok","Variabel","Resource","Elemen / Path FHIR","Nilai","Keterangan","Kode juga dipakai di","Muncul di (konsolidasi)","Status crosscheck","Catatan & aturan","Deskripsi variabel (sumber)","Induk","Peran","Level"]
W=[10,22,18,26,16,50,40,40,30,34,18,48,60,14,14,7]
for t,tname in M.TITLES:
    ws=wb.create_sheet(SHEET[t]); ws["A1"]=tname; ws["A1"].font=Font(name=F,bold=True,size=13)
    ws["A2"]="Sumber: "+M.SRC[t]; ws["A2"].font=Font(name=F,size=9,italic=True,color="555555")
    hdr(ws,4,H,W); row=5; prev=None
    REFT={}
    for r in M.REFS:
        if r["title"]==t: REFT.setdefault(r["tahap"],r)
    def refline(rr,tah):
        r=REFT.pop(tah,None)
        if not r: return rr
        txt="MENGIKUTI MODUL LAIN — "+r["note"]+(" Bagian yang dirujuk: "+r["isi"]+"." if r.get("isi") else "")+(" Kutipan playbook: “"+r["kutipan"]+"”" if r.get("kutipan") else "")+" Tautan: "+"; ".join(f'{x["judul"]} {x["url"]}' for x in r["targets"])
        cell(ws,rr,1,"(rujukan)",color=PL,bold=True); cell(ws,rr,2,tah,color="999999")
        cell(ws,rr,6,txt); ws.merge_cells(start_row=rr,start_column=6,end_row=rr,end_column=16)
        for i in range(1,17): ws.cell(rr,i).fill=PatternFill("solid",fgColor="F4E8ED"); ws.cell(rr,i).border=B
        return rr+1
    _order=[x for x in M.ALL if x["title"]==t]
    def _no(x):
        import re as _re; m=_re.match(r"^(\d{1,2})[.]",x or ""); return int(m.group(1)) if m else 999
    for n,v in enumerate(_order):
        if v["tahap"]!=prev:
            for tah in [k for k in list(REFT) if _no(k)<_no(v["tahap"])]:
                cell(ws,row,1,tah.upper(),bold=True,color=PL,size=10)
                for i in range(1,17): ws.cell(row,i).fill=STEP; ws.cell(row,i).border=B
                ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=16); row+=1
                row=refline(row,tah)
            cell(ws,row,1,v["tahap"].upper(),bold=True,color=PL,size=10)
            for i in range(1,17): ws.cell(row,i).fill=STEP; ws.cell(row,i).border=B
            ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=16); row+=1; prev=v["tahap"]
            row=refline(row,v["tahap"])
        if v.get("syn"):
            pos[v["id"]]=(SHEET[t],row)
            cell(ws,row,1,v["id"],bold=True,color=PL); cell(ws,row,4,"▼ "+v["var"],bold=True,color=PL)
            cell(ws,row,5,f'KELOMPOK — {len(v["anak"])} variabel di bawahnya'); cell(ws,row,15,"kelompok")
            for i in range(1,17): ws.cell(row,i).fill=PatternFill("solid",fgColor="F4E8ED"); ws.cell(row,i).border=B
            row+=1; continue
        g=M.GBY[v["gid"]]; pos[v["id"]]=(SHEET[t],row)
        others=[m["id"] for m in g["members"] if m["id"]!=v["id"]]
        muncul=f"{g['id']} {g['label']}"+(("\nJuga: "+"; ".join(vname(o) for o in others)) if others else "")+(("\nTerkait: "+"; ".join(f"{rid} {M.GBY[rid]['label']}" for rid,_ in g["related"][:6])) if g["related"] else "")
        note=" · ".join(x for x in [v["frek"] and "Frekuensi: "+v["frek"], v["cat"]] if x)
        for j,e in enumerate(v["elc"]):
            p,val,ket=e[:3]; lst=e[3] if len(e)>3 else None
            if lst: ket=f"→ Daftar {lst} ({len(M.LISTS[lst]['rows'])} baris)"
            np_=M.npath(p); al=[o for (pp,c),oo in v["also"].items() if pp==np_ and c==val for o in oo]
            vals=[v["id"],v["tahap"],v["kel"],("└ "*min(v.get("lvl",0),3))+v["var"],v["res"],p,val,ket,"; ".join(vname(o) for o in al[:4])+(" …" if len(al)>4 else ""),muncul if j==0 else None,g["status"] if j==0 else None,note if j==0 else None,(("[Playbook] " if M2.DESC[v["id"]]["src"]=="playbook" else "[Claude, belum diverifikasi] ")+M2.DESC[v["id"]]["t"]) if j==0 else None,
                  (v.get("ind") if j==0 else None),(v.get("per") if j==0 else None),(v.get("lvl",0) if j==0 else None)]
            for i,x in enumerate(vals,1):
                grey=(j>0 and i<=5)
                c=cell(ws,row,i,x,bold=(i==4 and j==0),color=("999999" if grey else ("3B2D5C" if i in (6,7) else "000000")),fill=BAND[n%2])
            if j==0: ws.cell(row,11).fill=PatternFill("solid",fgColor=STF[g["status"]]); links.append((ws.title,f"J{row}","Konsolidasi",g["id"]))
            if lst: links.append((ws.title,f"H{row}","LIST",lst))
            if al: links.append((ws.title,f"I{row}","VAR",al[0]))
            row+=1
        for x in v.get("sup",[]):
            if x["coded"]: nilai=" / ".join(f'{o[0]} | {o[1]} | {o[2]}' for o in x["body"][:12])
            else: nilai=" / ".join(str(o[0]) for o in x["body"][:12])
            if x["lists"]: nilai=(nilai+" ").strip()+" → Daftar "+", ".join(x["lists"])
            vals=[v["id"],v["tahap"],v["kel"],v["var"],v["res"],("*" if x["star"] else "")+x["key"],nilai or "(nilai tidak dirinci)",
                  "PELENGKAP dari playbook lain: "+", ".join(vname(i) for i in x["src"]),None,None,None,None,None]
            for i,xx in enumerate(vals,1):
                cell(ws,row,i,xx,color=("999999" if i<=5 else ("3B2D5C" if i in (6,7) else "000000")),fill=PatternFill("solid",fgColor="EAF1F7"))
            if x["lists"]: links.append((ws.title,f"H{row}","LIST",x["lists"][0]))
            row+=1
    for tah in list(REFT):
        cell(ws,row,1,tah.upper(),bold=True,color=PL,size=10)
        for i in range(1,17): ws.cell(row,i).fill=STEP; ws.cell(row,i).border=B
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=16); row+=1
        row=refline(row,tah)
    ws.freeze_panes="F5"; ws.auto_filter.ref=f"A4:P{row-1}"
# ---- Ringkasan
wr=wb.create_sheet("Ringkasan Crosscheck",1)
wr["A1"]="Ringkasan crosscheck: satu baris per konsep (deduplikasi lintas judul)"; wr["A1"].font=Font(name=F,bold=True,size=13)
HR=["ID konsep","Konsep","Status","Ada di (use case)","Variabel (use case · ID · nama)","Perbedaan (otomatis)","Catatan konsistensi","Konsep terkait"]
hdr(wr,3,HR,[10,30,20,18,44,80,50,40]); rr=4
for g in M.G:
    ms=g["members"]; ts=sorted({m["title"] for m in ms})
    def _w(ids): 
        by={}
        for i in ids: by.setdefault(M.BYID[i]["title"],[]).append(i)
        return ", ".join(t+(f" ×{len(v)}" if len(v)>1 else "") for t,v in by.items())
    dif="\n".join(("NILAI BEDA · " if x["kind"]=="beda" else "SEBAGIAN · ")+x["el"]+": "+" | ".join(f"{_w(ids)} = {t[:200]}" for t,ids in x["parts"]) for x in g["expl"])[:3000]
    rel="\n".join(f"{rid} {M.GBY[rid]['label']}"+(f" — {nt}" if nt else "") for rid,nt in g["related"])[:800]
    vals=[g["id"],g["label"],g["status"]," + ".join(t for t,_ in M.TITLES if t in ts),"\n".join(m["title"]+" · "+vname(m["id"]) for m in ms),dif,g["note"],rel]
    for i,x in enumerate(vals,1): cell(wr,rr,i,x,bold=(i==2))
    wr.cell(rr,3).fill=PatternFill("solid",fgColor=STF[g["status"]]); pos["R"+g["id"]]=rr
    links.append((wr.title,f"A{rr}","Konsolidasi",g["id"])); rr+=1
wr.freeze_panes="C4"; wr.auto_filter.ref=f"A3:H{rr-1}"
# ---- Konsolidasi
wk=wb.create_sheet("Konsolidasi",2)
wk["A1"]="Konsolidasi: elemen gabungan per konsep (baris unik) dengan use case tempat elemen itu muncul"; wk["A1"].font=Font(name=F,bold=True,size=13)
HK=["ID konsep","Konsep","Elemen / Path FHIR","Nilai (system | code | display)","Keterangan","Ada di (use case)","Dipakai oleh variabel","Perbedaan elemen"]
hdr(wk,3,HK,[10,28,46,60,36,22,30,16]); rk=4
for g in M.G:
    ms=g["members"]
    head=f"{g['id']} · {g['label']} — {g['status']}"
    cell(wk,rk,1,g["id"],bold=True,color=PL,fill=STEP,size=10); cell(wk,rk,2,g["label"],bold=True,color=PL,fill=STEP,size=10)
    cell(wk,rk,3,"Variabel: "+"; ".join(vname(m["id"]) for m in ms),fill=STEP); cell(wk,rk,4,g["status"],bold=True,fill=PatternFill("solid",fgColor=STF[g["status"]]))
    cell(wk,rk,5,g["note"] or None,fill=STEP)
    cell(wk,rk,6," + ".join(t for t,_ in M.TITLES if any(m["title"]==t for m in ms)),fill=STEP)
    cell(wk,rk,7,"← ringkasan",color="0563C1",fill=STEP); links.append((wk.title,f"G{rk}","Ringkasan Crosscheck","R"+g["id"]))
    links.append((wk.title,f"C{rk}","VAR",ms[0]["id"]))
    pos[g["id"]]=rk; rk+=1
    kind={x["el"]:x["kind"] for x in g["expl"]}
    for E in g["els"]:
        allm=len(E["ids"])==len(ms)
        items=[]
        if E["coded"]:
            for o in E["opts"].values(): items.append((f'{o["s"]} | {o["c"]} | {o["d"]}',o["k"],o["ids"],None))
        else:
            for x in E["vals"].values(): items.append((x["v"],x["k"],x["ids"],None))
        for l,li in E["lists"].items(): items.append((f"(daftar {l})","",li,l))
        if not items: items=[("","",E["ids"],None)]
        for j,(val,ket,ids,lst) in enumerate(items):
            ts={M.BYID[i]["title"] for i in ids}
            vals=[g["id"],g["label"],("*" if E["star"] else "")+E["key"] if j==0 else None,val,ket," + ".join(t for t,_ in M.TITLES if t in ts) if not (len(ids)==len(ms) and len(ms)>1) else "semua",", ".join(ids),({"beda":"nilai berbeda","sebagian":"hanya sebagian"}.get(kind.get(E["key"]),"") if j==0 else None)]
            for i,v in enumerate(vals,1): cell(wk,rk,i,v,color=("999999" if i<=2 else ("3B2D5C" if i in (3,4) else "000000")))
            if j==0 and kind.get(E["key"]): wk.cell(rk,8).fill=PatternFill("solid",fgColor="F4CCCC" if kind[E["key"]]=="beda" else "FCE8B2")
            if lst: links.append((wk.title,f"D{rk}","LIST",lst))
            rk+=1
wk.freeze_panes="C4"; wk.auto_filter.ref=f"A3:H{rk-1}"
# ---- Ringkasan Resource
wR=wb.create_sheet("Ringkasan Resource",3); wR["A1"]="Ringkasan Resource: elemen yang biasanya ada per resource, lintas semua playbook"; wR["A1"].font=Font(name=F,bold=True,size=13)
TL=[t for t,_ in M.TITLES]
hdr(wR,3,["Resource","Elemen","Tingkat","% variabel","Jumlah variabel","Jumlah induk","Jumlah use case","Wajib (*) di","Tipe data [x] (pilih salah satu)","Elemen baku FHIR?","Cakupan per use case (var dgn elemen / var resource)"],[20,26,10,10,10,10,10,30,40,12,90]); rR=4
LVF={"Inti":"D9EAD3","Umum":"FCE8B2","Kadang":"EEEEEE","Jarang":"EEEEEE"}
for rt,R in M2.RSOUT.items():
    cell(wR,rR,1,rt,bold=True,color=PL,fill=STEP,size=10); cell(wR,rR,2,f"{R['nv']} variabel ({R.get('npar')} induk) · {len(R['titles'])} use case",fill=STEP)
    for i in range(3,12): cell(wR,rR,i,None,fill=STEP)
    cell(wR,rR,11,("Belum pernah dipakai: "+", ".join(R["unused"])) if R["unused"] else None,fill=STEP); rR+=1
    for e in R["els"]:
        vals=[rt,e["el"],e["lvl"],e["pct"],e["n"],e.get("np"),len(e["titles"]),", ".join(e["star"]),"; ".join(f"{c} ({n})" for c,n in e["choices"].items()),"ya" if e["canon"] else ("tidak — cek" if R["canon"] else "?"),
              " · ".join(f"{t} {e['cov'][t]}/{R['tv'][t]}" for t in R["titles"] if e["cov"][t])]
        for i,v in enumerate(vals,1): cell(wR,rR,i,v,color=("3B2D5C" if i==2 else "000000"))
        wR.cell(rR,4).number_format="0%"; wR.cell(rR,3).fill=PatternFill("solid",fgColor=LVF[e["lvl"]]); rR+=1
    for x in R["stdonly"]:
        vals=[rt,x["path"],"—",None,0,0,0,None,None,"Lampiran Std. Terminologi",f"Hanya ada di Lampiran Standar Terminologi ({x['sec']}), {len(x['rows'])} kode; tidak dipakai di playbook mana pun"]
        for i,v in enumerate(vals,1): cell(wR,rR,i,v,color="7A2E4D")
        rR+=1
wR.freeze_panes="C4"; wR.auto_filter.ref=f"A3:K{rR-1}"
wN=wb.create_sheet("Resource - Nilai",4); wN["A1"]="Nilai / kode per path resource (playbook) + kode tambahan dari Lampiran Standar Terminologi SATUSEHAT v10.3"; wN["A1"].font=Font(name=F,bold=True,size=13)
hdr(wN,3,["Resource","Elemen","Path","system","code / nilai","display","Dipakai di (use case ×jumlah variabel)","Sumber"],[18,20,46,40,26,44,50,26]); rN=4
def _wt(ids):
    by={}
    for i in ids: by[M.BYID[i]["title"]]=by.get(M.BYID[i]["title"],0)+1
    return ", ".join(t+(f" ×{n}" if n>1 else "") for t,n in by.items())
STDF=PatternFill("solid",fgColor="F4E8ED")
for rt,R in M2.RSOUT.items():
    for e in R["els"]:
        for p in e["paths"]:
            if p["coded"]:
                for o in sorted(p["vals"],key=lambda o:-len(o[3])):
                    for i,v in enumerate([rt,e["el"],p["p"],o[0],o[1],o[2],_wt(o[3]),"Playbook"],1): cell(wN,rN,i,v,color=("3B2D5C" if i in (3,5) else "000000"))
                    rN+=1
                if p["std"]:
                    for o in p["std"]["extra"]:
                        for i,v in enumerate([rt,e["el"],p["p"],o[0],o[1],o[2] or o[3],None,"Lampiran Std. Terminologi (belum dipakai playbook)"],1): cell(wN,rN,i,v,fill=STDF)
                        rN+=1
            else:
                for o in sorted(p["vals"],key=lambda o:-len(o[1])):
                    for i,v in enumerate([rt,e["el"],p["p"],None,o[0],None,_wt(o[1]),"Playbook"],1): cell(wN,rN,i,v,color=("3B2D5C" if i in (3,5) else "000000"))
                    rN+=1
            for l,ids in p["lists"].items():
                for i,v in enumerate([rt,e["el"],p["p"],None,f"(daftar {l})",M.LISTS[l]["title"] if l in M.LISTS else "",_wt(ids),"Playbook"],1): cell(wN,rN,i,v)
                if l in M.LISTS: links.append((wN.title,f"E{rN}","LIST",l))
                rN+=1
    for x in R["stdonly"]:
        for o in x["rows"]:
            for i,v in enumerate([rt,x["el"],x["path"],o[0],o[1],o[2] or o[3],None,"Lampiran Std. Terminologi (path tidak dipakai playbook)"],1): cell(wN,rN,i,v,fill=STDF)
            rN+=1
wN.freeze_panes="D4"; wN.auto_filter.ref=f"A3:H{rN-1}"
# ---- Struktur induk-anak
wS=wb.create_sheet("Struktur Induk-Anak",5); wS["A1"]="Induk–anak: item kuesioner, komponen, bagian dokumen, dan kelompok playbook"; wS["A1"].font=Font(name=F,bold=True,size=13)
wS["A2"]=f'{HI.STAT["parents"]} induk ({HI.STAT["syn"]} kelompok sintetis) · {HI.STAT["child"]} variabel punya induk · peran: '+", ".join(f'{k} {v}' for k,v in HI.STAT["per"].items()); wS["A2"].font=Font(name=F,size=9,italic=True,color="555555")
hdr(wS,4,["ID konsep induk","Induk / sub-variabel","Peran","Use case","Konsep sub-variabel","Ada di use case"],[12,46,16,26,16,60]); rs_=5
for h in HI.H:
    cell(wS,rs_,1,h["id"],bold=True,color=PL,fill=STEP); cell(wS,rs_,2,h["label"],bold=True,fill=STEP)
    cell(wS,rs_,3,"konsep induk",fill=STEP); cell(wS,rs_,4," + ".join(h["titles"]),fill=STEP)
    cell(wS,rs_,5,None,fill=STEP); cell(wS,rs_,6,f'{len(h["m"])} induk · {h["nchild"]} sub-variabel',fill=STEP); rs_+=1
    for pid in h["m"]:
        p=M.BYID[pid]
        for i,x in enumerate([h["id"],"▼ "+p["var"],"induk ("+(p.get("per") or "induk")+")",p["title"],"",pid],1): cell(wS,rs_,i,x,color="7A2E4D")
        rs_+=1
    for r in h["rows"]:
        ada=" · ".join(f'{t}: {", ".join(ids)}' for t,ids in r["by"].items())
        for i,x in enumerate([h["id"],"    └ "+r["label"],"sub-variabel","",r["gid"] or "",ada],1): cell(wS,rs_,i,x)
        rs_+=1
wS.freeze_panes="C5"; wS.auto_filter.ref=f"A4:F{rs_-1}"
# ---- lists (single sheet)
wl=wb.create_sheet(LSHEET); wl["A1"]="Daftar pilihan jawaban & lampiran (system – code – display – keterangan)"; wl["A1"].font=Font(name=F,bold=True,size=13)
hdr(wl,3,["ID daftar","Judul daftar","Dipakai oleh variabel","Kolom 1","Kolom 2","Kolom 3","Kolom 4","Kolom 5","Kolom 6"],[10,40,30,34,16,40,40,20,20]); rl=4; LPOS={}
for k,Ld in M.LISTS.items():
    LPOS[k]=rl
    cell(wl,rl,1,k,bold=True,color=PL,fill=STEP,size=10); cell(wl,rl,2,Ld["title"],bold=True,fill=STEP); cell(wl,rl,3,", ".join(Ld.get("users",[])),fill=STEP)
    for i,c in enumerate(Ld["cols"],4): cell(wl,rl,i,c,bold=True,fill=STEP)
    rl+=1
    for rw in Ld["rows"]:
        cell(wl,rl,1,k,color="999999"); cell(wl,rl,2,None); cell(wl,rl,3,None)
        for i,v in enumerate(rw,4): cell(wl,rl,i,v,color=("3B2D5C" if i==5 else "000000"))
        rl+=1
wl.freeze_panes="D4"; wl.auto_filter.ref=f"A3:I{rl-1}"
# ---- apply links
for sh,ref,tsh,tgt in links:
    if tsh=="VAR": s,rw=pos[tgt]; loc=f"'{s}'!A{rw}"
    elif tsh=="Konsolidasi": loc=f"'Konsolidasi'!A{pos[tgt]}"
    elif tsh=="Ringkasan Crosscheck": loc=f"'Ringkasan Crosscheck'!A{pos[tgt]}"
    elif tsh=="LIST": loc=f"'{LSHEET}'!A{LPOS[tgt]}"
    else: loc=f"'{tsh}'!A{tgt}"
    c=wb[sh][ref]; c.hyperlink=Hyperlink(ref=ref,location=loc); c.font=Font(name=F,size=9,color="0563C1",underline="single",bold=c.font.bold)
out="/mnt/user-data/outputs/SATUSEHAT_Semua_UseCase_Crosscheck.xlsx"; wb.save(out); print(out, {ws.title:ws.max_row for ws in wb})
